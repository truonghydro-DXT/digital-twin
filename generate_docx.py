from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Bao_cao_thuat_toan_phuong_phap_ky_thuat_digital_twin.docx"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(result.rows[0].cells[index], header, True, (255, 255, 255))
        shade(result.rows[0].cells[index], "1F4E78")
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if len(result.rows) % 2 == 0:
                shade(cells[index], "EAF2F8")
    if widths:
        for row in result.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return result


def bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(40, 40, 40)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
for name, size, color in (("Title", 22, (31, 78, 121)), ("Heading 1", 16, (31, 78, 121)), ("Heading 2", 12.5, (46, 116, 145))):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(*color)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("BÁO CÁO THUẬT TOÁN, PHƯƠNG PHÁP\nVÀ KỸ THUẬT ỨNG DỤNG DIGITAL TWIN")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Phạm vi: D:\\WORK\\digital-twin | Ngày lập: " + date.today().strftime("%d/%m/%Y"))
run.italic = True
run.font.color.rgb = RGBColor(89, 89, 89)

doc.add_paragraph(
    "Tài liệu này mô tả phiên bản ứng dụng hiện có trong thư mục digital-twin, "
    "bao gồm pipeline photogrammetry, cơ chế sử dụng Agisoft Metashape Professional, "
    "xử lý dữ liệu DJI L2, dịch vụ viewer và các sản phẩm đầu ra. Mô tả được xây dựng "
    "từ README và mã nguồn tại thời điểm lập tài liệu."
)

doc.add_heading("1. Tổng quan và mục tiêu", level=1)
doc.add_paragraph(
    "Ứng dụng tạo một digital twin 3D từ chuỗi ảnh UAV DJI L2 có dữ liệu định vị, "
    "sau đó xuất mây điểm, mô hình mesh, DEM, orthomosaic và báo cáo. Một viewer web "
    "độc lập cho phép kiểm tra các sản phẩm trên nền bản đồ, xem dữ liệu 3D và nạp shapefile."
)
table(doc, ["Thành phần", "Vai trò", "Công nghệ chính"], [
    ("digital-twin.py", "Điều phối pipeline photogrammetry và xuất sản phẩm", "Python, Metashape API"),
    ("UAV-digital-twin/src", "Các lớp project, ảnh, point cloud, mesh, monitor", "Python"),
    ("view-output.py", "Web viewer, quản lý job DTW/Tiles và đường dẫn dữ liệu", "Flask, JavaScript"),
    ("lib/dem-3d", "Hiển thị DEM/orthomosaic 3D và điểm", "Three.js"),
    ("lib/static/vendors", "Bản đồ và hiển thị 3D", "Leaflet, Cesium, MapLibre, Plotly"),
])

doc.add_heading("2. Dữ liệu đầu vào và tiền điều kiện", level=1)
bullet(doc, "Ảnh UAV: các tệp *_D.JPG hoặc JPEG trong thư mục LiDAR-image.")
bullet(doc, "Tệp .MRK: bản ghi timestamp, latitude, longitude, ellipsoidal height và vận tốc VN/VE/VV.")
bullet(doc, "Các tệp .LDR/.LDRT là dữ liệu LiDAR thô DJI; pipeline hiện không đọc trực tiếp. Muốn tích hợp LiDAR cần chuyển trước sang LAS/LAZ bằng DJI Terra.")
bullet(doc, "Python 3.11, wheel Metashape phù hợp và Agisoft Metashape Professional đã kích hoạt.")
bullet(doc, "Đối với desktop, script có thể relaunch qua metashape.exe để dùng license GUI; viewer truyền process bằng đúng Python interpreter của nó.")

doc.add_heading("3. Kiến trúc xử lý", level=1)
doc.add_paragraph("Luồng xử lý tổng quát:")
for item in [
    "Quét và lập danh mục ảnh theo biểu thức nhận dạng chỉ số ảnh.",
    "Đọc MRK, ghép bản ghi định vị với ảnh theo chỉ số hoặc thứ tự tương ứng.",
    "Khởi tạo Metashape Document/Chunk, thêm ảnh và gán tọa độ camera.",
    "Lọc ảnh, matching keypoints, align camera và tối ưu camera.",
    "Tính depth map, tạo dense point cloud và xuất LAS.",
    "Dựng mesh, texture, DEM và orthomosaic.",
    "Ghi project, catalog, camera reference, log, report và summary.",
]:
    numbered(doc, item)

doc.add_heading("4. Thuật toán và phương pháp", level=1)
doc.add_heading("4.1. Nhận dạng ảnh và ghép GPS", level=2)
doc.add_paragraph(
    "Ảnh được chọn theo phần mở rộng JPG/JPEG và chỉ số bốn chữ số trước hậu tố _D. "
    "Mỗi dòng MRK được chuyển thành bản ghi có index và tọa độ. Hàm attach_gps gán lat, "
    "lon và ellh vào PhotoRecord; ảnh không ghép được GPS vẫn được ghi nhận để phục vụ kiểm tra."
)
doc.add_heading("4.2. Photogrammetry từ ảnh", level=2)
doc.add_paragraph(
    "Metashape thực hiện matching các đặc trưng ảnh, align camera dựa trên các điểm tương đồng "
    "và bundle adjustment để ước lượng đồng thời vị trí, hướng camera cùng tham số liên quan. "
    "Các bước này tạo sparse tie points và nền tảng cho tính toán bề mặt dày đặc."
)
doc.add_heading("4.3. Tái tạo bề mặt dày đặc", level=2)
doc.add_paragraph(
    "Depth map được ước lượng từ các camera đã align. Depth map được chuyển thành dense cloud, "
    "sau đó xuất sang LAS. Đây là bước có chi phí tính toán lớn nhất và nhạy với chất lượng ảnh, "
    "độ chồng phủ, ánh sáng, texture bề mặt và độ chính xác camera."
)
doc.add_heading("4.4. Mesh, texture, DEM và orthomosaic", level=2)
doc.add_paragraph(
    "Dense cloud hoặc depth data được dùng để dựng mesh và texture. DEM biểu diễn bề mặt độ cao "
    "dạng raster; orthomosaic là ảnh trực giao đã hiệu chỉnh hình học. Viewer có thể cắt DEM theo "
    "bounding box của orthomosaic để đồng bộ phạm vi hiển thị."
)

doc.add_heading("5. Kỹ thuật tính toán CPU/GPU và độ tin cậy", level=1)
doc.add_paragraph(
    "Tham số --gpu-mask là bitmask thiết bị. Device 0 tương ứng bit 0 và device 1 tương ứng bit 1. "
    "Trong giao diện Run DTW, chế độ GPU hiện truyền --gpu-mask 10, còn CPU truyền --cpu-only. "
    "CPU vẫn được bật để xử lý các công đoạn không hỗ trợ GPU."
)
table(doc, ["Chế độ", "Tham số", "Mục đích"], [
    ("GPU device 0", "--gpu-mask 1", "Dùng GPU đầu tiên; cần kiểm tra thứ tự thiết bị."),
    ("GPU device 1", "--gpu-mask 10", "Dùng GPU thứ hai, thường để tránh Intel HD trên máy lai."),
    ("CPU", "--cpu-only", "Tắt GPU, phù hợp khi OpenCL lỗi hoặc cần đường chạy ổn định."),
])
doc.add_paragraph("Các preset chất lượng ảnh hưởng trực tiếp đến độ phân giải depth map và thời gian:")
table(doc, ["Preset", "Độ phân giải tương đối depth", "Sử dụng"], [
    ("low", "x8", "Kiểm tra nhanh, máy yếu hoặc thử pipeline."),
    ("medium", "x4", "Cân bằng chất lượng và thời gian; mặc định."),
    ("high", "x2", "Máy mạnh, RAM lớn, cần chất lượng cao."),
])
doc.add_paragraph(
    "Dấu hiệu xác nhận GPU là log của depth map có [GPU]. Nếu chỉ có [CPU 1], GPU chưa được sử dụng. "
    "Nếu tiến trình dừng lâu tại preparing cameras images hoặc tie points loaded, cần thử mask khác "
    "hoặc chuyển CPU; không nên gọi enumGPUDevices thường xuyên vì có thể treo trên Intel HD."
)

doc.add_heading("6. Phân tích dữ liệu và kiểm soát chất lượng", level=1)
doc.add_heading("6.1. Kiểm tra đầu vào", level=2)
bullet(doc, "Đếm số ảnh hợp lệ và số bản ghi MRK; kiểm tra ảnh có chỉ số tương ứng.")
bullet(doc, "Kiểm tra phạm vi lat/lon, ellh và tính liên tục của đường bay.")
bullet(doc, "Xác nhận ảnh đủ chồng phủ và không bị mờ quá mức trước khi matching.")
doc.add_heading("6.2. Kiểm tra trung gian", level=2)
bullet(doc, "Theo dõi số camera đã align, tie points và log tiến độ.")
bullet(doc, "So sánh thời gian depth map giữa CPU và GPU ở cùng preset.")
bullet(doc, "Kiểm tra bounding box địa lý của DEM/orthomosaic và tính nhất quán CRS.")
doc.add_heading("6.3. Kiểm tra đầu ra", level=2)
bullet(doc, "LAS có point count lớn hơn 0 và hiển thị đúng phạm vi.")
bullet(doc, "OBJ có vertex/face và texture liên kết đúng.")
bullet(doc, "DEM/orthomosaic mở được, có CRS và bounds hợp lệ.")
bullet(doc, "twin_summary.json, camera_reference.csv và run.log phản ánh đúng phiên chạy.")

doc.add_heading("7. Ứng dụng viewer và kỹ thuật hiển thị", level=1)
doc.add_paragraph(
    "view-output.py chạy Flask server. Viewer đọc sản phẩm từ output-new, hiển thị dashboard đường bay "
    "bằng Plotly, raster 2D bằng Leaflet, DEM/điểm bằng Three.js và 3D Tiles bằng Cesium. Các endpoint "
    "cung cấp metadata, file media, tile và trạng thái job."
)
bullet(doc, "Run DTW tạo subprocess với working directory của ứng dụng, truyền images/output và chế độ CPU/GPU.")
bullet(doc, "Run Tiles tạo tileset từ sản phẩm đã xuất.")
bullet(doc, "Job có polling trạng thái, phần trăm tiến độ và log giới hạn trong bộ nhớ viewer.")
bullet(doc, "Shapefile được chuyển sang GeoJSON EPSG:4326 để nạp lên bản đồ; trường thuộc tính có thể dùng để dựng khối 3D.")
bullet(doc, "LAS được stream dạng nhị phân và lấy mẫu để hiển thị hiệu quả trên trình duyệt.")

doc.add_heading("8. Các tệp đầu ra", level=1)
table(doc, ["Tệp/thư mục", "Nội dung", "Cách sử dụng"], [
    ("digital_twin.psx", "Project Metashape", "Mở lại, tiếp tục hoặc xuất lại sản phẩm."),
    ("point_cloud/point_cloud.las", "Mây điểm LAS", "Phân tích điểm, GIS, CAD hoặc viewer 3D."),
    ("model/model.obj", "Mesh 3D", "Mô hình bề mặt dạng OBJ."),
    ("texture/texture.jpg", "Ảnh texture", "Phủ vật liệu cho mesh."),
    ("dem/dem.tif", "Mô hình số độ cao", "Phân tích địa hình và hiển thị DEM."),
    ("orthomosaic/orthomosaic.tif", "Ảnh trực giao", "Bản đồ nền và kiểm tra hình học."),
    ("report.pdf", "Báo cáo Metashape", "Tóm tắt dự án và thông tin xử lý."),
    ("mission_catalog.json", "Danh mục ảnh/MRK", "Truy vết dữ liệu nguồn."),
    ("camera_reference.csv", "Tọa độ camera", "Vẽ đường bay và kiểm tra GPS."),
    ("run.log", "Log UTF-8 theo thời gian", "Chẩn đoán phiên xử lý."),
    ("twin_summary.json", "Tóm tắt phiên chạy", "Metadata phiên bản, quality, camera và project."),
])

doc.add_heading("9. Vận hành đề xuất", level=1)
code(doc, r"cd D:\WORK\digital-twin")
code(doc, r".\venv\Scripts\python.exe digital-twin.py --gpu-mask 10 --quality medium")
code(doc, r".\venv\Scripts\python.exe view-output.py")
doc.add_paragraph("Quy trình chạy an toàn:")
for item in [
    "Chạy --list-only hoặc kiểm tra dữ liệu để xác nhận ảnh và MRK.",
    "Chạy thử quality low nếu chưa biết độ ổn định của dữ liệu.",
    "Chạy medium bằng GPU đã xác định; kiểm tra log có [GPU].",
    "Không dùng --reset khi project đã align và cần tiếp tục; dùng --resume.",
    "Sau khi hoàn tất, mở viewer và kiểm tra lần lượt overview, orthomosaic, DEM, LAS và mesh.",
]:
    numbered(doc, item)

doc.add_heading("10. Giới hạn và rủi ro kỹ thuật", level=1)
bullet(doc, "License của wheel Python và license GUI Metashape có thể khác nhau; cần để script relaunch qua Pro trên desktop hoặc kích hoạt wheel riêng trên môi trường không GUI.")
bullet(doc, "GPU mask phụ thuộc thứ tự thiết bị do Metashape nhận diện; mask 1 không phải lúc nào cũng là GPU rời.")
bullet(doc, "Dữ liệu LDR/LDRT chưa được đưa trực tiếp vào pipeline, nên kết quả hiện là photogrammetry từ ảnh và GPS.")
bullet(doc, "Chất lượng hình học phụ thuộc mạnh vào chồng phủ ảnh, blur, ánh sáng, calibration và độ tin cậy MRK.")
bullet(doc, "Viewer lấy mẫu dữ liệu lớn để hiển thị; số điểm nhìn thấy trên web có thể nhỏ hơn point count gốc.")

doc.add_heading("11. Kết luận", level=1)
doc.add_paragraph(
    "Ứng dụng hiện thực hóa một pipeline digital twin khép kín từ ảnh UAV: chuẩn hóa dữ liệu, "
    "định vị camera, tái tạo 3D bằng Metashape, xuất dữ liệu chuẩn GIS/3D và kiểm tra trực quan "
    "trên web. Thiết kế tách phần xử lý nặng khỏi viewer, hỗ trợ tiếp tục project và lựa chọn CPU/GPU, "
    "phù hợp cho quy trình thử nghiệm đến sản xuất. Khi cần độ tin cậy cao hơn, nên bổ sung thống kê "
    "điểm kiểm tra mặt đất, sai số camera, báo cáo reprojection error và đối soát với LAS độc lập."
)

doc.save(OUTPUT)
print(OUTPUT)